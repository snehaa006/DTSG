#!/usr/bin/env python3
"""Build dtsg_internship_report.docx from the LaTeX source.

The report is written in LaTeX; this produces the Word version that the
department asks for.  Pandoc does the conversion, but three things in the
source need help first, and the resulting docx needs a little layout work:

  * the two TikZ diagrams are cropped out of the built PDF as images, since
    pandoc cannot draw TikZ;
  * cross-references, citations and caption numbers are resolved here,
    because pandoc has no LaTeX counters;
  * a few displays (\\underbrace, cases, \\mathtt tokens) come out mangled in
    Word's math format, so they are rewritten as plain lines;
  * afterwards the docx gets A4 pages, sane table column widths, black
    headings, a title page, page numbers and a filled contents page.

Requires: pandoc, libreoffice (for the contents page numbers), and the
Python packages pypdfium2, Pillow and python-docx.

    python3 make_docx.py
"""

import re
import shutil
import subprocess
import sys
from pathlib import Path

import docx
import pypdfium2 as pdfium
from PIL import Image
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

HERE = Path(__file__).resolve().parent
TEX = HERE / 'dtsg_internship_report.tex'
PDF = HERE / 'dtsg_internship_report.pdf'
OUT = HERE / 'dtsg_internship_report.docx'
BUILD = HERE / 'build'

TEXT_CM = 21.0 - 2 * 2.54          # A4 width less one-inch margins
W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


# --------------------------------------------------------------- figures ---
def crop_figures():
    """Crop the two TikZ figures out of the built PDF at 300 dpi."""
    pdf = pdfium.PdfDocument(str(PDF))
    boxes = [(15, (280, 745, 2210, 1645), 'fig1.png'),    # architecture
             (18, (290, 2395, 2200, 2665), 'fig2.png')]   # supersede chain
    for page, box, name in boxes:
        pdf[page].render(scale=300 / 72).to_pil().crop(box).save(BUILD / name)


# ------------------------------------------------------------ latex prep ---
ALGORITHM = r"""\begin{quote}
\textbf{Algorithm 1: DTSG ingest --- from a message to a consistent state graph}
\end{quote}
\begin{verbatim}
 1: procedure INGEST(u: user, m: message)
 2:   e <- InsertEvent(u, m)        # own txn; commits before derived state
 3:   R <- ExtractFacts(m)          # LLM, schema-enforced triples
 4:   F <- Normalize(R)             # folding, snake_case, synonym map
 5:   for each fact f in F do
 6:     v <- Embed(f)               # null on failure; degrade, do not drop
 7:     C <- ExactMatch(u, f) U VectorMatch(u, v)   # sim. floor 0.55
 8:     if C = {} then
 9:       (r, T) <- (additive, {})  # no candidates, no possible conflict
10:     else
11:       (r, T) <- Classify(f, C)  # LLM; targets reconciled against C
12:     end if
13:     if r = reinforce then
14:       BumpConfidence(T_1)       # also stamps last_reinforced_at
15:     else
16:       begin transaction
17:       n <- InsertMemory(u, f, v, e)
18:       if r = supersede then
19:         Expire(T, n)            # sets superseded_by; guarded on ACTIVE
20:       end if
21:       commit
22:     end if
23:   end for
24: end procedure
\end{verbatim}
"""

EQ_SCORE = r"""\begin{quote}
\noindent (1)\quad $\mathrm{score}(f, q, t) \;=\; \max\!\big(0,\ \cos(\mathbf{e}_f, \mathbf{e}_q)\big) \;\times\; w(f, t) \;\times\; e^{-\lambda\,\Delta(f,t)}$

\noindent \emph{is it relevant} $\times$ \emph{is it still true} $\times$ \emph{how stale is it}
\end{quote}
"""

EQ_WEIGHT = r"""\begin{quote}
\noindent (2)\quad $w(f,t) = w_{\text{exp}}$ \quad if \texttt{valid\_from}$(f) > t$ (it had not happened yet)

\noindent \quad $w(f,t) = 1.0$ \quad if \texttt{valid\_until}$(f)$ is \textsc{null} or \texttt{valid\_until}$(f) > t$

\noindent \quad $w(f,t) = w_{\text{exp}}$ \quad otherwise (it had already stopped being true)
\end{quote}
"""

RESOLUTIONS = r"""\begin{quote}
\noindent \texttt{user speaks English} + \texttt{user speaks Hindi} $\rightarrow$ \textsc{additive} (multi-valued)

\noindent \texttt{user lives\_in Delhi} + \texttt{user lives\_in Berlin} $\rightarrow$ \textsc{supersede} (single-valued)
\end{quote}
"""

INGEST_PATH_OLD = ('\\[\n\\text{message} \\rightarrow \\text{event (committed)} '
                   '\\rightarrow \\text{facts}\n\\rightarrow \\text{candidates} '
                   '\\rightarrow \\text{classification}\n\\rightarrow \\text{writes}.\n\\]')
INGEST_PATH_NEW = ('\\begin{quote}\n\\noindent message $\\rightarrow$ event (committed) '
                   '$\\rightarrow$ facts $\\rightarrow$ candidates $\\rightarrow$ '
                   'classification $\\rightarrow$ writes.\n\\end{quote}')

# tabularx X-columns carry no width pandoc can use; spell them out
COLUMN_SPECS = {
    r'{l l X}':                   r'{p{3.0cm} p{3.0cm} p{10.0cm}}',
    r'{p{2.6cm} X p{3.5cm}}':     r'{p{2.6cm} p{9.9cm} p{3.5cm}}',
    r'{p{2.6cm} p{3.2cm} X}':     r'{p{2.6cm} p{3.2cm} p{10.2cm}}',
    r'{p{3.1cm} X p{3.4cm}}':     r'{p{3.1cm} p{9.5cm} p{3.4cm}}',
    r'{p{3.4cm} p{4.6cm} c c X}': r'{p{3.4cm} p{4.6cm} p{1.7cm} p{1.7cm} p{4.6cm}}',
    r'{X l l c}':                 r'{p{5.2cm} p{4.0cm} p{4.4cm} p{2.4cm}}',
}


def resolve_labels(src):
    """Map every \\label to the number LaTeX would print for it."""
    labels = {}
    sec = sub = subsub = 0
    tab = fig = eq = 0
    current = ''
    stack = []
    token = re.compile(
        r'\\(section|subsection|subsubsection)(\*?)\{|'
        r'\\begin\{(table|figure|equation|longtable)\*?\}|'
        r'\\end\{(table|figure|equation|longtable)\*?\}|'
        r'\\label\{([^}]*)\}')
    for m in token.finditer(src):
        head, star, begin, end, label = m.groups()
        if head:
            if star:
                continue
            if head == 'section':
                sec, sub, subsub = sec + 1, 0, 0
                current = str(sec)
            elif head == 'subsection':
                sub, subsub = sub + 1, 0
                current = f'{sec}.{sub}'
            else:
                subsub += 1
                current = f'{sec}.{sub}.{subsub}'
        elif begin:
            if begin in ('table', 'longtable'):
                tab += 1
            elif begin == 'figure':
                fig += 1
            elif begin == 'equation':
                eq += 1
            stack.append(begin)
        elif end:
            if stack:
                stack.pop()
        elif label is not None:
            env = stack[-1] if stack else None
            if env in ('table', 'longtable'):
                labels[label] = str(tab)
            elif env == 'figure':
                labels[label] = str(fig)
            elif env == 'equation':
                labels[label] = str(eq)
            else:
                labels[label] = current
    return labels


def number_captions(src):
    """Prefix each caption with 'Table n: ' / 'Figure n: '."""
    out, idx, tables, figures = [], 0, 0, 0
    pattern = re.compile(r'\\begin\{(table|figure|longtable)\*?\}')
    while True:
        m = pattern.search(src, idx)
        if not m:
            out.append(src[idx:])
            return ''.join(out)
        kind = m.group(1)
        if kind == 'figure':
            figures += 1
            prefix = 'Figure %d: ' % figures
        else:
            tables += 1
            prefix = 'Table %d: ' % tables
        caption = src.find(r'\caption{', m.end())
        env_end = src.find(r'\end{%s}' % kind, m.end())
        out.append(src[idx:m.end()])
        if caption != -1 and caption < env_end:
            out.append(src[m.end():caption + len(r'\caption{')])
            out.append(prefix)
            idx = caption + len(r'\caption{')
        else:
            idx = m.end()


def swap_environment(src, begin, end, replacement, marker):
    """Replace the environment containing `marker` with `replacement`."""
    i = src.index(marker)
    a = src.rindex(begin, 0, i)
    b = src.index(end, i) + len(end)
    return src[:a] + replacement + src[b:]


def prepare_tex():
    src = TEX.read_text()

    # TikZ pictures -> the cropped images
    for image in ('fig1.png', 'fig2.png'):
        start = src.index(r'\begin{tikzpicture}')
        stop = src.index(r'\end{tikzpicture}') + len(r'\end{tikzpicture}')
        before, after = src[:start], src[stop:]
        wrapper = re.search(r'\\resizebox\{[^}]*\}\{[^}]*\}\{%?\s*$', before)
        if wrapper:                       # drop the \resizebox around it
            before = before[:wrapper.start()]
            after = after.lstrip()[1:]
        src = (before + '\\includegraphics[width=\\textwidth]{%s}\n' % image + after)

    # algpseudocode -> a plain listing
    start = src.index(r'\begin{algorithm}')
    stop = src.index(r'\end{algorithm}') + len(r'\end{algorithm}')
    src = src[:start] + ALGORITHM + src[stop:]
    src = src.replace(r'Algorithm~\ref{alg:ingest}', 'Algorithm 1')

    labels = resolve_labels(src)
    src = re.sub(r'\\eqref\{([^}]*)\}', lambda m: '(%s)' % labels.get(m.group(1), '?'), src)
    src = re.sub(r'\\ref\{([^}]*)\}', lambda m: labels.get(m.group(1), '?'), src)

    keys = re.findall(r'\\bibitem\{([^}]*)\}', src)
    numbers = {key: i + 1 for i, key in enumerate(keys)}
    src = re.sub(r'\\cite\{([^}]*)\}',
                 lambda m: '[' + ', '.join(str(numbers.get(k.strip(), '?'))
                                           for k in m.group(1).split(',')) + ']', src)

    src = number_captions(src)

    # math inside a caption does not survive the docx writer
    def caption_text(m):
        body = m.group(0).replace(r'$\to$', '\u2192').replace(r'$\times$', '\u00d7')
        return re.sub(r'\$([A-Za-z0-9\\{}.,\- ]+)\$',
                      lambda t: t.group(1).replace('\\', ''), body)
    src = re.sub(r'\\caption\{(?:[^{}]|\{[^{}]*\})*\}', caption_text, src)

    # \paragraph headings read better run-in, as they do in the PDF
    src = re.sub(r'\\paragraph\{((?:[^{}]|\{[^{}]*\})*)\}\s*',
                 lambda m: r'\noindent\textbf{%s} ' % m.group(1), src)

    # bibliography: a real heading and visible [n] markers
    src = src.replace('\\begin{thebibliography}{99}\\small', '\\section*{References}\n')
    src = src.replace('\\end{thebibliography}', '')
    src = re.sub(r'\\bibitem\{([^}]*)\}',
                 lambda m: '\n\n{[}%d{]} ' % numbers[m.group(1)], src)

    for old, new in COLUMN_SPECS.items():
        src = src.replace(r'\begin{tabularx}{\textwidth}' + old, r'\begin{tabular}' + new)
    src = src.replace(r'\end{tabularx}', r'\end{tabular}')

    src = swap_environment(src, r'\begin{equation}', r'\end{equation}',
                           EQ_SCORE, r'\underbrace')
    src = swap_environment(src, r'\begin{equation}', r'\end{equation}',
                           EQ_WEIGHT, r'\begin{cases}')
    src = swap_environment(src, r'\begin{align*}', r'\end{align*}',
                           RESOLUTIONS, r'\textsc{additive} \quad')
    src = src.replace(INGEST_PATH_OLD, INGEST_PATH_NEW)

    (BUILD / 'report.tex').write_text(src)


def run_pandoc():
    subprocess.run(
        ['pandoc', 'report.tex', '-f', 'latex', '-t', 'docx',
         '--number-sections', '--toc', '--toc-depth=3', '-o', str(OUT)],
        cwd=BUILD, check=True)


# ------------------------------------------------------------ docx layout ---
# column widths per table, in the order the tables appear in the document
TABLE_WIDTHS = [
    [3.2, 3.6, 9.1],                       # 1  contributions
    [2.6, 9.8, 3.5],                       # 2  comparative analysis
    [2.6, 3.2, 10.1],                      # 3  other related work
    [12.4, 3.5],                           # 4  dataset description
    [3.1, 9.4, 3.4],                       # 5  data attributes
    [3.4, 4.6, 1.7, 1.7, 4.5],             # 6  benchmark cases
    [2.6, 13.3],                           # -  resolution set
    [3.4, 8.6, 3.9],                       # -  retrieval modes
    [5.0, 3.7, 4.4, 2.8],                  # 7  per-case outcome
    [6.9, 4.5, 4.5],                       # 8  aggregate
    [4.4, 1.9, 1.8, 2.0, 2.3, 2.0, 0.8],   # 9  term ablation
]


def set_column_widths(table, widths):
    widths = [w * TEXT_CM / sum(widths) for w in widths]
    table.autofit = False
    layout = table._tbl.tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        table._tbl.tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for col, cm in zip(grid.findall(qn('w:gridCol')), widths):
            col.set(qn('w:w'), str(int(Cm(cm).twips)))
    for row in table.rows:
        for cell, cm in zip(row.cells, widths):
            cell.width = Cm(cm)


def page_break_after(paragraph):
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def page_number_footer(document):
    for section in document.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.text = ''
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        begin = OxmlElement('w:fldChar')
        begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = ' PAGE '
        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')
        for element in (begin, instr, end):
            run._r.append(element)


def lay_out_docx():
    d = docx.Document(str(OUT))

    for section in d.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        for margin in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            setattr(section, margin, Cm(2.54))

    for table, widths in zip(d.tables, TABLE_WIDTHS):
        if len(table.columns) == len(widths):
            set_column_widths(table, widths)
        else:
            print('table column count changed; widths not applied', file=sys.stderr)

    for shape in d.inline_shapes:                     # keep figures inside the margins
        if shape.width > Emu(int(Cm(TEXT_CM).emu)):
            ratio = shape.height / shape.width
            shape.width = Cm(TEXT_CM)
            shape.height = Cm(TEXT_CM * ratio)

    for name in ('Source Code', 'Verbatim Char'):
        try:
            d.styles[name].font.size = Pt(8.5)
        except KeyError:
            pass

    # front matter: title page, then one page per preliminary section
    ps = d.paragraphs
    for i in list(range(0, 8)) + [8, 13, 14, 16, 19, 20, 22]:
        ps[i].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i in (8, 14, 16, 20):                          # the preliminary headings
        for run in ps[i].runs:
            run.bold = True
    ps[0].runs[0].font.size = Pt(20)
    for run in ps[1].runs:
        run.font.size = Pt(14)
    for i in (6, 7, 13, 19, 22):
        page_break_after(ps[i])

    # the contents belongs after the front matter, not before the title page
    body = d.element.body
    sdt = body.find(qn('w:sdt'))
    if sdt is not None:
        body.remove(sdt)
        ps[23]._p.addprevious(sdt)
        ps[23]._p.addprevious(parse_xml(
            f'<w:p {W_NS}><w:r><w:br w:type="page"/></w:r></w:p>'))

    page_number_footer(d)
    d.save(str(OUT))

    # heading colours are theme references python-docx cannot override
    import zipfile
    tmp = OUT.with_suffix('.docx.tmp')
    with zipfile.ZipFile(OUT) as zin, \
            zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'word/styles.xml':
                xml = data.decode('utf-8')
                xml = re.sub(r'<w:color[^>]*w:themeColor[^>]*/>',
                             '<w:color w:val="000000"/>', xml)
                data = xml.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, OUT)


def fill_contents():
    """Give the TOC field a cached result, so the page is not blank in viewers
    that never evaluate fields.  Word still refreshes it on update."""
    # render into build/ so this never touches the LaTeX-built PDF next door
    shutil.copy(OUT, BUILD / OUT.name)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', OUT.name],
                   cwd=BUILD, check=True, capture_output=True)
    rendered = (BUILD / OUT.name).with_suffix('.pdf')
    pdf = pdfium.PdfDocument(str(rendered))
    pages = [re.sub(r'\s+', '', pdf[i].get_textpage().get_text_range())
             for i in range(len(pdf))]

    d = docx.Document(str(OUT))
    rows = []
    for p in d.paragraphs:
        if not p.style.name.startswith('Heading') or not p.text.strip():
            continue
        level = int(p.style.name[-1])
        if level > 3:
            continue
        text = p.text.replace('\t', '  ')
        needle = re.sub(r'\s+', '', text)
        page = next((i + 1 for i, body in enumerate(pages) if needle in body), None)
        if page is None:
            continue
        escaped = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        rows.append(
            f'<w:p {W_NS}><w:pPr><w:pStyle w:val="TOC{level}"/>'
            f'<w:ind w:left="{(level - 1) * 360}"/>'
            f'<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9020"/></w:tabs>'
            f'</w:pPr><w:r><w:t xml:space="preserve">{escaped}</w:t></w:r>'
            f'<w:r><w:tab/><w:t>{page}</w:t></w:r></w:p>')

    sdt = d.element.body.find(qn('w:sdt'))
    content = sdt.find(qn('w:sdtContent'))
    field = content.findall(qn('w:p'))[-1]
    field.addprevious(parse_xml(
        f'<w:p {W_NS}><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/>'
        f'<w:instrText xml:space="preserve">TOC \\o "1-3" \\h \\z \\u</w:instrText>'
        f'<w:fldChar w:fldCharType="separate"/></w:r></w:p>'))
    for row in rows:
        field.addprevious(parse_xml(row))
    field.addprevious(parse_xml(
        f'<w:p {W_NS}><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'))
    content.remove(field)
    d.save(str(OUT))
    return len(rows)


def main():
    BUILD.mkdir(exist_ok=True)
    crop_figures()
    prepare_tex()
    run_pandoc()
    lay_out_docx()
    entries = fill_contents()
    print(f'{OUT.relative_to(HERE.parent)} written ({entries} contents entries)')


if __name__ == '__main__':
    main()
